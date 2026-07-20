from pathlib import Path
from urllib.parse import urlencode

from pptx import Presentation
from pptx.util import Inches
from playwright.sync_api import sync_playwright


EDGE_PATH = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
SLIDE_WIDTH = 1920
SLIDE_HEIGHT = 1080


def export_html_deck_to_pptx(source_url, slide_count, output_path, work_dir):
    """Render each source HTML slide in Edge and place the rendered page into PPTX."""
    output_path = Path(output_path)
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)

    if not EDGE_PATH.exists():
        raise RuntimeError(f"Microsoft Edge not found: {EDGE_PATH}")

    images = []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            headless=True,
            executable_path=str(EDGE_PATH),
            args=["--force-color-profile=srgb"],
        )
        context = browser.new_context(
            viewport={"width": SLIDE_WIDTH, "height": SLIDE_HEIGHT},
            device_scale_factor=1,
        )
        page = context.new_page()
        for index in range(1, slide_count + 1):
            query = urlencode({"preview": index, "export": "pptx"})
            page.goto(f"{source_url}?{query}", wait_until="networkidle")
            page.wait_for_timeout(700)
            image_path = work_dir / f"slide-{index:02d}.png"
            page.screenshot(path=str(image_path), type="png")
            images.append(image_path)
        context.close()
        browser.close()

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    for image_path in images:
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(image_path), 0, 0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
    presentation.save(str(output_path))
